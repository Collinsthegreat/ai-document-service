"""
Service for extracting text from PDF and DOCX documents.
"""
import io
from typing import BinaryIO

import PyPDF2
import pdfplumber
from docx import Document as DocxDocument

from app.core.exceptions import PDFExtractionError, DocxExtractionError
from app.core.logging import get_logger

logger = get_logger(__name__)


class TextExtractionService:
    """
    Service for extracting text from various document formats.
    
    Supports PDF and DOCX files with fallback extraction methods.
    """
    
    @staticmethod
    async def extract_from_pdf(file_content: bytes, filename: str) -> str:
        """
        Extract text from PDF file with fallback mechanism.
        
        First attempts extraction using pdfplumber (more reliable),
        falls back to PyPDF2 if that fails.
        
        Args:
            file_content: Binary content of PDF file.
            filename: Original filename (for error reporting).
            
        Returns:
            Extracted text content.
            
        Raises:
            PDFExtractionError: If text extraction fails completely.
        """
        text_parts = []
        
        try:
         
            logger.debug(f"Attempting PDF extraction with pdfplumber: {filename}")
            
            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                for page_num, page in enumerate(pdf.pages, start=1):
                    try:
                        page_text = page.extract_text()
                        if page_text:
                            text_parts.append(page_text)
                    except Exception as e:
                        logger.warning(
                            f"Failed to extract page {page_num} with pdfplumber",
                            extra={"file_name": filename, "page": page_num, "error": str(e)}
                        )
            
            if text_parts:
                extracted_text = "\n\n".join(text_parts)
                logger.info(
                    f"PDF extraction successful using pdfplumber",
                    extra={
                        "file_name": filename,
                        "pages": len(text_parts),
                        "text_length": len(extracted_text)
                    }
                )
                return extracted_text
                
        except Exception as e:
            logger.warning(
                f"pdfplumber extraction failed, trying PyPDF2",
                extra={"file_name": filename, "error": str(e)}
            )
        
        try:
           
            logger.debug(f"Attempting PDF extraction with PyPDF2: {filename}")
            
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
            
            for page_num in range(len(pdf_reader.pages)):
                try:
                    page = pdf_reader.pages[page_num]
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
                except Exception as e:
                    logger.warning(
                        f"Failed to extract page {page_num + 1} with PyPDF2",
                        extra={"file_name": filename, "page": page_num + 1, "error": str(e)}
                    )
            
            if text_parts:
                extracted_text = "\n\n".join(text_parts)
                logger.info(
                    f"PDF extraction successful using PyPDF2",
                    extra={
                        "file_name": filename,
                        "pages": len(text_parts),
                        "text_length": len(extracted_text)
                    }
                )
                return extracted_text
            
        except Exception as e:
            logger.error(
                f"Both PDF extraction methods failed",
                extra={"file_name": filename, "error": str(e)},
                exc_info=True
            )
            raise PDFExtractionError(str(e), filename)
        
      
        error_msg = "No text content found in PDF"
        logger.error(error_msg, extra={"file_name": filename})
        raise PDFExtractionError(error_msg, filename)
    
    @staticmethod
    async def extract_from_docx(file_content: bytes, filename: str) -> str:
        """
        Extract text from DOCX file.
        
        Args:
            file_content: Binary content of DOCX file.
            filename: Original filename (for error reporting).
            
        Returns:
            Extracted text content.
            
        Raises:
            DocxExtractionError: If text extraction fails.
        """
        try:
            logger.debug(f"Attempting DOCX extraction: {filename}")
            
            
            doc = DocxDocument(io.BytesIO(file_content))
            
            
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
          
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)
            
            if not text_parts:
                error_msg = "No text content found in DOCX"
                logger.error(error_msg, extra={"file_name": filename})
                raise DocxExtractionError(error_msg, filename)
            
            extracted_text = "\n\n".join(text_parts)
            
            logger.info(
                f"DOCX extraction successful",
                extra={
                    "file_name": filename,
                    "paragraphs": len(doc.paragraphs),
                    "tables": len(doc.tables),
                    "text_length": len(extracted_text)
                }
            )
            
            return extracted_text
            
        except DocxExtractionError:
            raise
        except Exception as e:
            logger.error(
                f"DOCX extraction failed",
                extra={"file_name": filename, "error": str(e)},
                exc_info=True
            )
            raise DocxExtractionError(str(e), filename)
    
    async def extract_text(
        self,
        file_content: bytes,
        filename: str,
        file_extension: str
    ) -> str:
        """
        Extract text from document based on file type.
        
        Args:
            file_content: Binary content of file.
            filename: Original filename.
            file_extension: File extension (pdf, docx).
            
        Returns:
            Extracted text content.
            
        Raises:
            PDFExtractionError: If PDF extraction fails.
            DocxExtractionError: If DOCX extraction fails.
            ValueError: If file extension is not supported.
        """
        extension = file_extension.lower()
        
        if extension == "pdf":
            return await self.extract_from_pdf(file_content, filename)
        elif extension == "docx":
            return await self.extract_from_docx(file_content, filename)
        else:
            raise ValueError(f"Unsupported file extension: {extension}")